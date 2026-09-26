"""Create paired, fictional PDF/DOCX originals for the document comparator."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


PDF_FIELDS = [
    ("Titular de la propiedad", "Ana María Torres"),
    ("Número de identificación", "52.123.456"),
    ("Matrícula inmobiliaria", "050N-1234567"),
    ("Nombre del predio", "La Aurora"),
    ("Municipio", "Chía"),
    ("Departamento", "Cundinamarca"),
    ("Dirección", "Vereda Fagua, Lote 7"),
    ("Área del predio", "12,50 hectáreas"),
    ("Fecha de adquisición", "15/03/2024"),
    ("Servidumbre inscrita", "No registra"),
]

DOCX_FIELDS = [
    ("Titular de la propiedad", "Ana María Torres"),
    ("Número de identificación", "91.987.654"),
    ("Matrícula inmobiliaria", "050N-1234567"),
    ("Nombre del predio", "El Roble"),
    ("Municipio", "Chía"),
    ("Departamento", "Cundinamarca"),
    ("Dirección", "Vereda Fagua, Lote 8"),
    ("Área del predio", "12,55 hectáreas"),
    ("Fecha de adquisición", "15/03/2024"),
]


def make_pdf(destination: Path) -> None:
    pdfmetrics.registerFont(TTFont("ArialFixture", "C:/Windows/Fonts/arial.ttf"))
    pdfmetrics.registerFont(TTFont("ArialFixtureBold", "C:/Windows/Fonts/arialbd.ttf"))
    destination.parent.mkdir(parents=True, exist_ok=True)
    page = canvas.Canvas(str(destination), pagesize=letter)
    width, height = letter
    x = 62
    page.setTitle("Ficha de título predial prueba A")
    page.setAuthor("Territorium")
    page.setFillColor(HexColor("#50615A"))
    page.setFont("ArialFixtureBold", 9)
    page.drawString(x, height - 61, "TERRITORIUM  /  PRUEBA DE COMPARACIÓN")
    page.setFillColor(HexColor("#172820"))
    page.setFont("ArialFixtureBold", 21)
    page.drawString(x, height - 98, "Ficha de título predial")
    page.setFillColor(HexColor("#4C5C53"))
    page.setFont("ArialFixture", 10)
    page.drawString(x, height - 119, "Referencia de prueba A  |  Predio La Aurora")
    page.drawString(x, height - 140, "Datos declarados para el cotejo con la ficha registral B.")
    page.setStrokeColor(HexColor("#CBD6CE"))
    page.line(x, height - 160, width - x, height - 160)
    y = height - 190
    for label, value in PDF_FIELDS:
        page.setFont("ArialFixtureBold", 10)
        page.setFillColor(HexColor("#35463B"))
        page.drawString(x, y, f"{label}:")
        page.setFont("ArialFixture", 10)
        page.setFillColor(HexColor("#172820"))
        page.drawString(x + 188, y, value)
        y -= 39
    page.setStrokeColor(HexColor("#CBD6CE"))
    page.line(x, y + 8, width - x, y + 8)
    page.setFont("ArialFixture", 8)
    page.setFillColor(HexColor("#68776F"))
    page.drawString(x, y - 15, "Documento ficticio para pruebas. Sin validez jurídica ni datos personales reales.")
    page.showPage()
    page.save()


def make_docx(destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(23, 40, 32)
    normal.paragraph_format.space_after = Pt(9)
    title_style = doc.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(7)

    eyebrow = doc.add_paragraph("TERRITORIUM  /  PRUEBA DE COMPARACIÓN")
    eyebrow.style = "Normal"
    eyebrow.paragraph_format.space_after = Pt(12)
    eyebrow.runs[0].font.size = Pt(8)
    eyebrow.runs[0].font.bold = True
    title = doc.add_paragraph("Ficha registral predial", style="Title")
    # Word's built-in Title style can add a colored bottom border.
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "nil")
    borders.append(bottom)
    title._p.get_or_add_pPr().append(borders)
    subtitle = doc.add_paragraph("Referencia de prueba B  |  Predio El Roble")
    subtitle.paragraph_format.space_after = Pt(13)
    doc.add_paragraph("Datos declarados para el cotejo con la ficha de título predial A.")
    for label, value in DOCX_FIELDS:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(12)
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(value)
    footer = doc.add_paragraph("Documento ficticio para pruebas. Sin validez jurídica ni datos personales reales.")
    footer.paragraph_format.space_before = Pt(12)
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(104, 119, 111)
    doc.core_properties.title = "Ficha registral predial prueba B"
    doc.core_properties.author = "Territorium"
    doc.save(destination)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("format", choices=["pdf", "docx"])
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    {"pdf": make_pdf, "docx": make_docx}[args.format](args.destination)
