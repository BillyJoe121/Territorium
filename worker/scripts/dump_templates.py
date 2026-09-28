import os
import json
import docx

templates_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "plantillas documentos finales"))

files = {
    "escritura": "ESCRITURA TOL-ANZ-045.docx",
    "linderos": "ID02 descripción de linderos.docx",
    "minuta": "MINUTA_TIPO_TERRITORIUM.docx"
}

out_dir = os.path.join(os.path.dirname(__file__), "extracted_templates")
os.makedirs(out_dir, exist_ok=True)

for key, filename in files.items():
    filepath = os.path.join(templates_dir, filename)
    doc = docx.Document(filepath)
    
    doc_data = {
        "filename": filename,
        "paragraphs": [],
        "tables": []
    }
    
    for i, p in enumerate(doc.paragraphs):
        doc_data["paragraphs"].append({
            "index": i,
            "text": p.text,
            "style": p.style.name if p.style else None
        })
        
    for t_idx, table in enumerate(doc.tables):
        t_data = []
        for r in table.rows:
            r_data = [c.text.strip() for c in r.cells]
            t_data.append(r_data)
        doc_data["tables"].append({
            "index": t_idx,
            "rows": t_data
        })
        
    out_file = os.path.join(out_dir, f"{key}.json")
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(doc_data, f, ensure_ascii=False, indent=2)
        
    # Also write a readable text file
    txt_file = os.path.join(out_dir, f"{key}.txt")
    with open(txt_file, "w", encoding="utf-8") as f:
        f.write(f"=== {filename} ===\n\n")
        for p in doc_data["paragraphs"]:
            if p["text"].strip():
                f.write(p["text"] + "\n\n")
        if doc_data["tables"]:
            f.write("\n=== TABLES ===\n")
            for t in doc_data["tables"]:
                f.write(f"\n--- Table {t['index']} ---\n")
                for r in t["rows"]:
                    f.write(" | ".join(r) + "\n")
                    
    print(f"Exported {key}: {len(doc_data['paragraphs'])} paragraphs, {len(doc_data['tables'])} tables.")
