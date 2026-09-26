import io
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock

import docx

from app.document_comparison import compare_documents, parse_original, process_comparison_job, validate_comparison


def make_document(name: str, text: str):
    document = docx.Document()
    document.add_paragraph(text)
    stream = io.BytesIO()
    document.save(stream)
    return parse_original(stream.getvalue(), name, name)


def side(document, value):
    fragment = document.fragments[0]
    return {'fragment_id': fragment.fragment_id, 'quote': fragment.text, 'value': value}


class DocumentComparisonTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.left = make_document('a.docx', 'Titular: María Elena Rojas')
        self.right = make_document('b.docx', 'Titular: María Elena Rojas')

    def test_exact_match_requires_identical_literal_values_and_verified_quotes(self):
        result = validate_comparison({'fields': [{
            'key': 'titular', 'label': 'Titular',
            'left': side(self.left, 'María Elena Rojas'),
            'right': side(self.right, 'María Elena Rojas'),
        }]}, self.left, self.right)
        self.assertEqual(result['fields'][0]['status'], 'exact')
        self.assertEqual(result['fields'][0]['left']['location'], 'Párrafo 1')
        self.assertEqual(result['counts']['exact'], 1)

    def test_small_character_difference_is_not_exact(self):
        different = make_document('b.docx', 'Titular: Maria Elena Rojas')
        result = validate_comparison({'fields': [{
            'key': 'titular', 'label': 'Titular',
            'left': side(self.left, 'María Elena Rojas'),
            'right': side(different, 'Maria Elena Rojas'),
        }]}, self.left, different)
        self.assertEqual(result['fields'][0]['status'], 'near')

    def test_missing_evidence_is_red_and_never_gains_a_fake_location(self):
        result = validate_comparison({'fields': [{
            'key': 'titular', 'label': 'Titular',
            'left': side(self.left, 'María Elena Rojas'), 'right': None,
        }]}, self.left, self.right)
        self.assertEqual(result['fields'][0]['status'], 'different')
        self.assertIsNone(result['fields'][0]['right'])

    def test_significantly_different_values_are_red(self):
        different = make_document('b.docx', 'Titular: Asociación Río Claro')
        result = validate_comparison({'fields': [{
            'key': 'titular', 'label': 'Titular',
            'left': side(self.left, 'María Elena Rojas'),
            'right': side(different, 'Asociación Río Claro'),
        }]}, self.left, different)
        self.assertEqual(result['fields'][0]['status'], 'different')

    def test_hallucinated_quote_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'NO_VERIFIABLE_FIELDS'):
            validate_comparison({'fields': [{
                'key': 'titular', 'label': 'Titular',
                'left': {'fragment_id': self.left.fragments[0].fragment_id, 'quote': 'Otro nombre', 'value': 'Otro'},
                'right': None,
            }]}, self.left, self.right)

    async def test_ai_response_uses_existing_provider_contract(self):
        payload = {'fields': [{
            'key': 'titular', 'label': 'Titular',
            'left': side(self.left, 'María Elena Rojas'),
            'right': side(self.right, 'María Elena Rojas'),
        }]}
        import json
        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=AsyncMock(
            return_value=SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(payload)))])
        ))))
        result = await compare_documents(client, 'test-model', self.left, self.right)
        self.assertEqual(result['counts']['exact'], 1)
        self.assertEqual(client.chat.completions.create.await_args.kwargs['response_format'], {'type': 'json_object'})

    async def test_worker_persists_only_a_verified_result(self):
        import json
        payload = {'fields': [{
            'key': 'titular', 'label': 'Titular',
            'left': side(self.left, 'María Elena Rojas'),
            'right': side(self.right, 'María Elena Rojas'),
        }]}
        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=AsyncMock(
            return_value=SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(payload)))])
        ))))
        source = docx.Document()
        source.add_paragraph('Titular: María Elena Rojas')
        stream = io.BytesIO()
        source.save(stream)
        gateway = SimpleNamespace(
            get_comparison_document=AsyncMock(side_effect=[
                {'id': 'a.docx', 'original_name': 'a.docx', 'storage_path': 'left'},
                {'id': 'b.docx', 'original_name': 'b.docx', 'storage_path': 'right'},
            ]),
            download=AsyncMock(return_value=stream.getvalue()),
            finish_comparison_job=AsyncMock(return_value=True),
        )
        await process_comparison_job(gateway, {
            'id': 'job', 'project_id': 'project', 'left_document_id': 'a.docx',
            'right_document_id': 'b.docx', 'lease_token': 'lease',
        }, client, 'test-model')
        saved = gateway.finish_comparison_job.await_args.args
        self.assertEqual(saved[2]['counts']['exact'], 1)
        self.assertIsNone(saved[3])


if __name__ == '__main__':
    unittest.main()
