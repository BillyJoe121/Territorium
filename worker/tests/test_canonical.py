import unittest
import docx
import io

from app.preprocessing.docx_parser import parse_docx
from app.preprocessing.pdf_parser import parse_pdf
from app.models.canonical import ScanClassification, UnitType


class CanonicalParserTests(unittest.TestCase):
    def test_docx_parser_preserves_structure_and_locators(self) -> None:
        doc = docx.Document()
        doc.add_paragraph("ESTUDIO DE TÍTULOS — PREDIO LA ESPERANZA")
        doc.add_paragraph("Folio de matrícula inmobiliaria: 050N-204581")
        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "Propietario"
        tbl.cell(0, 1).text = "Documento"
        tbl.cell(1, 0).text = "María Elena Rojas"
        tbl.cell(1, 1).text = "43123456"

        stream = io.BytesIO()
        doc.save(stream)
        file_bytes = stream.getvalue()

        canonical = parse_docx(file_bytes=file_bytes, original_name="test_estudio.docx")

        self.assertEqual(canonical.original_name, "test_estudio.docx")
        self.assertEqual(canonical.overall_scan_status, ScanClassification.TEXTUAL)
        self.assertTrue(len(canonical.sha256) == 64)
        self.assertGreaterEqual(len(canonical.fragments), 3)

        # Check locators
        p1 = canonical.fragments[0]
        self.assertEqual(p1.locator.unit_type, UnitType.PARAGRAPH)
        self.assertIn("ESTUDIO DE TÍTULOS", p1.text)

        # Check table fragment
        table_frag = [f for f in canonical.fragments if f.locator.unit_type == UnitType.TABLE][0]
        self.assertIsNotNone(table_frag.table_data)
        self.assertEqual(table_frag.table_data[1][0], "María Elena Rojas")

    def test_pdf_parser_empty_raises_error(self) -> None:
        with self.assertRaises(ValueError):
            parse_pdf(b"", "empty.pdf")


if __name__ == "__main__":
    unittest.main()
