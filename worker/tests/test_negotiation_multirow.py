"""La plantilla de negociación trae un predio por fila: se extraen todos, con su FMI."""
import io
import unittest

import openpyxl

from app.extractors.negotiation_extractor import NegotiationExtractor
from app.utils.spanish_currency import amount_in_words


def _workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CONSOLIDADO"
    ws.append(["DISEÑO", "CARPETA", "CODIGO", "FMI", "ÁREA PREDIO (m2)", "Valor oferta No. 1 (50%)", "Valor oferta No. 2 (80%)", "Valor oferta No. 3 (100%)", "VALOR NEGOCIADO"])
    ws.append(["", "", "", "", "", "", "", "", ""])  # subencabezado vacío
    ws.append(["SUPLENCIA", "TOL-ANZ-103", "73043000200060049000", "350-129805", 20000, 1485141.94, 2376227.11, 2970283.89, 1800000])
    ws.append(["SUPLENCIA", "TOL-ANZ-045", "73043000200020024000", "350-108418", 69524, 7628711.51, 12205938.42, 15257423.03, 9628712])
    ws.append(["SUPLENCIA", "TOL-ALV-029", "73026000200070001000", "Sin información", 187500, 7829759.64, 12527615.43, 15659519.28, None])
    # Hoja resumen que repite una carpeta sin ofertas: no debe duplicar la fila.
    resumen = wb.create_sheet("RESUMEN")
    resumen.append(["CARPETA", "FMI", "Valor oferta No. 1"])
    resumen.append(["TOL-ANZ-103", "350-129805", None])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class NegotiationMultiRowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.payload = NegotiationExtractor().extract_from_xlsx_bytes(_workbook(), "tabla.xlsx", "doc-1")

    def test_extracts_one_row_per_property_with_fmi_and_cadastral_id(self) -> None:
        rows = {r["property_code"]: r for r in self.payload.negotiations}
        self.assertEqual(set(rows), {"TOL-ANZ-103", "TOL-ANZ-045", "TOL-ALV-029"})
        self.assertEqual(rows["TOL-ANZ-103"]["fmi"], "350-129805")
        self.assertEqual(rows["TOL-ANZ-103"]["cadastral_id"], "73043000200060049000")
        self.assertEqual(rows["TOL-ALV-029"]["fmi"], "")  # "Sin información" no es un FMI

    def test_rounds_offers_and_writes_missing_letters(self) -> None:
        row = next(r for r in self.payload.negotiations if r["property_code"] == "TOL-ANZ-103")
        self.assertEqual(row["first_offer_numbers"], "$ 1.485.142")
        self.assertEqual(row["first_offer_letters"], "UN MILLÓN CUATROCIENTOS OCHENTA Y CINCO MIL CIENTO CUARENTA Y DOS PESOS")
        self.assertTrue(row["letters_generated"])
        self.assertEqual(row["values_match"], "Sí, coinciden")

    def test_negotiated_value_column_is_not_taken_as_an_offer(self) -> None:
        row = next(r for r in self.payload.negotiations if r["property_code"] == "TOL-ANZ-045")
        self.assertEqual(row["third_offer_numbers"], "$ 15.257.423")


class FormulaWithoutCachedValueTests(unittest.TestCase):
    def test_rows_with_uncalculated_formulas_are_reported_not_dropped(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["CARPETA", "FMI", "VALOR TOTAL", "Valor oferta No. 1", "Valor oferta No. 2"])
        ws.append(["TOL-ANZ-103", "350-129805", 2970283.89, "=C2*0.5", "=C2*0.8"])
        out = io.BytesIO()
        wb.save(out)  # openpyxl guarda las fórmulas sin su valor calculado
        payload = NegotiationExtractor().extract_from_xlsx_bytes(out.getvalue(), "sin_calcular.xlsx", "doc-1")
        row = payload.negotiations[0]
        self.assertEqual(row["fmi"], "350-129805")
        self.assertIn("Abre el archivo en Excel", row["values_match"])
        self.assertTrue(any("fórmulas sin valor calculado" in d for d in payload.discrepancies))


class OfferComparisonTests(unittest.TestCase):
    def test_apocope_forms_are_accepted(self) -> None:
        from app.utils.spanish_currency import compare_number_and_letters
        self.assertTrue(compare_number_and_letters(24_321_971, amount_in_words(24_321_971))[0])
        self.assertTrue(compare_number_and_letters(1_000_000, "UN MILLÓN DE PESOS")[0])
        self.assertFalse(compare_number_and_letters(24_321_971, "VEINTE MILLONES DE PESOS")[0])


class AmountInWordsTests(unittest.TestCase):
    def test_canonical_spanish(self) -> None:
        self.assertEqual(amount_in_words(1_000_000), "UN MILLÓN DE PESOS")
        self.assertEqual(amount_in_words(21_000), "VEINTIÚN MIL PESOS")
        self.assertEqual(amount_in_words(16_316), "DIECISÉIS MIL TRESCIENTOS DIECISÉIS PESOS")
        self.assertEqual(amount_in_words(9_628_712), "NUEVE MILLONES SEISCIENTOS VEINTIOCHO MIL SETECIENTOS DOCE PESOS")


if __name__ == "__main__":
    unittest.main()
