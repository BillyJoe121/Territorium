import io
import unittest
import openpyxl

from app.processing.consolidator import (
    ConsolidatedMasterRecord,
    consolidate_property_records,
)
from app.export.excel_generator import generate_consolidated_excel_bytes


class TestPhase5ConsolidationExcel(unittest.TestCase):
    def setUp(self):
        self.titles_payload = {
            "folio": "050N-204581",
            "cadastral_id": "05001010400230012000",
            "property_name": "La Esperanza",
            "municipality": "Rionegro",
            "department": "Antioquia",
            "village": "La Esperanza",
            "owners": [{"name": "María Elena Rojas", "document_type": "CC", "document_number": "43.123.456"}],
            "acquisition_mode": "Compraventa",
            "boundaries": "Norte: El Roble; Sur: quebrada",
            "boundaries_document": "Escritura 1240",
            "legal_conditions": "Sin anotaciones restrictivas",
            "justice_ministry_case": "MJ-2026-001",
            "urt_case": "URT-2026-001",
            "urt_territorial_direction": "Antioquia",
        }

        self.plans_payload = {
            "plans": [
                {
                    "plan_name": "Plano Servidumbre Tramo 12",
                    "easement_area_numbers": "4580",
                    "easement_length_numbers": "458",
                    "easement_width_numbers": "10",
                    "infrastructure_count_numbers": "8",
                    "plan_scale": "1:2000",
                    "voltage_level": "230 kV",
                }
            ]
        }

        self.negotiation_payload = {
            "property_code": "PREDIO-050N",
            "first_offer_numbers": "$ 218.450.000",
            "second_offer_numbers": "$ 232.800.000",
            "third_offer_numbers": "—",
            "values_match": "Sí, coinciden",
        }

    def test_consolidate_property_records_deterministic(self):
        record = consolidate_property_records(
            titles_payload=self.titles_payload,
            titles_version_id="t-v1",
            plans_payload=self.plans_payload,
            plans_version_id="p-v1",
            negotiation_payload=self.negotiation_payload,
            negotiation_version_id="n-v1",
            user_id="user-approver",
        )

        self.assertIsInstance(record, ConsolidatedMasterRecord)
        self.assertEqual(record.folio, "050N-204581")
        self.assertEqual(record.cadastral_id, "05001010400230012000")
        self.assertEqual(record.property_name, "La Esperanza")
        self.assertEqual(record.owners, "María Elena Rojas (CC 43.123.456)")
        self.assertEqual(record.easement_area, "4580")
        self.assertEqual(record.easement_length, "458")
        self.assertEqual(record.first_offer, "$ 218.450.000")
        self.assertEqual(record.second_offer, "$ 232.800.000")
        self.assertEqual(record.values_match, "Sí, coinciden")

        self.assertEqual(record.metadata["titles_result_version_id"], "t-v1")
        self.assertEqual(record.metadata["plans_result_version_id"], "p-v1")
        self.assertEqual(record.metadata["negotiation_result_version_id"], "n-v1")
        self.assertEqual(record.metadata["consolidated_by"], "user-approver")

    def test_generate_consolidated_excel_openpyxl(self):
        record = consolidate_property_records(
            titles_payload=self.titles_payload,
            titles_version_id="t-v1",
            plans_payload=self.plans_payload,
            plans_version_id="p-v1",
            negotiation_payload=self.negotiation_payload,
            negotiation_version_id="n-v1",
            user_id="user-approver",
        )

        excel_bytes = generate_consolidated_excel_bytes(
            record=record,
            project_id="project-xyz",
            project_name="Interconexión 230kV",
            version_number=1,
        )

        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 1000)

        # Parse generated workbook
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertEqual(len(wb.sheetnames), 2)
        self.assertIn("Consolidado Predial", wb.sheetnames)
        self.assertIn("Metadatos y Trazabilidad", wb.sheetnames)

        # Verify sheet 1 columns
        ws1 = wb["Consolidado Predial"]
        self.assertEqual(ws1.cell(row=1, column=1).value, "CARPETA")
        self.assertEqual(ws1.cell(row=1, column=2).value, "FOLIO DE MATRICULA")
        self.assertEqual(ws1.cell(row=2, column=2).value, "050N-204581")
        self.assertEqual(ws1.cell(row=2, column=5).value, "La Esperanza")

        # Verify sheet 2 metadata
        ws2 = wb["Metadatos y Trazabilidad"]
        self.assertEqual(ws2.cell(row=1, column=1).value, "Propiedad / Metadato")
        self.assertEqual(ws2.cell(row=2, column=2).value, "Interconexión 230kV")
        self.assertEqual(ws2.cell(row=3, column=2).value, "project-xyz")
        self.assertEqual(ws2.cell(row=4, column=2).value, "v1")


if __name__ == "__main__":
    unittest.main()
