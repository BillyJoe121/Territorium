import unittest
import openpyxl
import io

from app.preprocessing.negotiation_reader import read_negotiation_xlsx
from app.extractors.negotiation_extractor import NegotiationExtractor


class NegotiationReaderTests(unittest.TestCase):
    def test_read_negotiation_xlsx_extracts_offers_and_coordinates(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ofertas"

        # Headers
        ws.cell(row=2, column=1, value="CARPETA")
        ws.cell(row=2, column=2, value="Valor oferta No. 1 (Números)")
        ws.cell(row=2, column=3, value="Valor oferta No. 1 (Letras)")
        ws.cell(row=2, column=4, value="Valor oferta No. 2 (Números)")
        ws.cell(row=2, column=5, value="Valor oferta No. 2 (Letras)")

        # Row 1
        ws.cell(row=3, column=1, value="TOL-ANZ-103")
        ws.cell(row=3, column=2, value=1485142)
        ws.cell(row=3, column=3, value="UN MILLON CUATROCIENTOS OCHENTA Y CINCO MIL CIENTO CUARENTA Y DOS")
        ws.cell(row=3, column=4, value=2376227)
        ws.cell(row=3, column=5, value="DOS MILLONES TRESCIENTOS SETENTA Y SEIS MIL DOSCIENTOS VEINTISIETE")

        stream = io.BytesIO()
        wb.save(stream)
        file_bytes = stream.getvalue()

        result = read_negotiation_xlsx(
            file_bytes=file_bytes,
            original_name="negociacion.xlsx",
            document_id="doc-neg-1",
            target_property_code="TOL-ANZ-103",
        )

        self.assertEqual(len(result.rows), 1)
        row = result.active_row
        self.assertIsNotNone(row)
        self.assertEqual(row.property_code, "TOL-ANZ-103")
        self.assertEqual(row.first_offer_number, 1485142.0)
        self.assertIn("B3", row.first_offer_number_cell or "")

        # Now test with extractor
        extractor = NegotiationExtractor()
        extracted = extractor.extract_from_xlsx_bytes(
            file_bytes=file_bytes,
            original_name="negociacion.xlsx",
            document_id="doc-neg-1",
            target_property_code="TOL-ANZ-103",
        )

        self.assertTrue(extracted.first_offer_matches)
        self.assertTrue(extracted.second_offer_matches)
        self.assertEqual(extracted.values_match, "Sí, coinciden")
        self.assertEqual(len(extracted.discrepancies), 0)

    def test_correspondencia_format_with_duplicate_headers_and_predio_columns(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Hoja1"

        # Headers mimicking CORRESPONDENCIA.xlsx
        ws.cell(row=1, column=1, value="CARPETA")
        ws.cell(row=1, column=2, value="FOLIO DE MATRICULA")
        ws.cell(row=1, column=9, value="NOMBRE DEL PREDIO")
        ws.cell(row=1, column=18, value="DOCUMENTO QUE CONTIENE LOS LINDEROS DEL PREDIO")
        ws.cell(row=1, column=51, value="VALOR OFERTA 1 (NUMEROS)")
        ws.cell(row=1, column=52, value="VALOR OFERTA 1 (LETRAS)")
        ws.cell(row=1, column=53, value="VALOR OFERTA 2 (NUMEROS)")
        ws.cell(row=1, column=54, value="VALOR OFERTA 2 (LETRAS)")
        ws.cell(row=1, column=55, value="VALOR OFERTA 3 (NUMEROS)")
        ws.cell(row=1, column=56, value="VALOR OFERTA 3 (NUMEROS)")  # Duplicate/mislabeled header

        # Row 2 (Data)
        ws.cell(row=2, column=1, value="TOL-ANZ-045")
        ws.cell(row=2, column=2, value="350-108418")
        ws.cell(row=2, column=9, value="LA PLAYA")
        ws.cell(row=2, column=18, value="resolución 694 del 30 de junio de 1994...")
        ws.cell(row=2, column=51, value=7628711.0)
        ws.cell(row=2, column=52, value="SIETE MILLONES SEISCIENTOS VEINTIOCHO MIL SETECIENTOS ONCE PESOS")
        ws.cell(row=2, column=53, value=12205938.0)
        ws.cell(row=2, column=54, value="DOCE MILLONES DOSCIENTOS CINCO MIL NOVECIENTOS TREINTA Y OCHO PESOS")
        ws.cell(row=2, column=55, value=15257423.0)
        ws.cell(row=2, column=56, value="QUINCE MILLONES DOSCIENTOS CINCUENTA Y SIETE MIL CUATROCIENTOS VEINTITRÉS PESOS")

        stream = io.BytesIO()
        wb.save(stream)
        file_bytes = stream.getvalue()

        extractor = NegotiationExtractor()
        extracted = extractor.extract_from_xlsx_bytes(
            file_bytes=file_bytes,
            original_name="correspondencia.xlsx",
            document_id="doc-corr-1",
        )

        self.assertEqual(extracted.property_code, "TOL-ANZ-045")
        self.assertEqual(extracted.first_offer_numeric_value, 7628711.0)
        self.assertEqual(extracted.second_offer_numeric_value, 12205938.0)
        self.assertEqual(extracted.third_offer_numeric_value, 15257423.0)
        self.assertEqual(extracted.third_offer_letters, "QUINCE MILLONES DOSCIENTOS CINCUENTA Y SIETE MIL CUATROCIENTOS VEINTITRÉS PESOS")
        self.assertEqual(extracted.values_match, "Sí, coinciden")
        self.assertEqual(len(extracted.discrepancies), 0)


if __name__ == "__main__":
    unittest.main()
